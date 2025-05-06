from appJar import gui
from geometry_package.shapes import Rectangle, Triangle, Trapezoid
from docx import Document

class SimpleGeometryCalculator:
    def __init__(self):
        self.app = gui("Фигуры", "400x300")
        self.app.addButton("Прямоугольник", lambda: self.calculate("Прямоугольник"))
        self.app.addButton("Треугольник", lambda: self.calculate("Треугольник"))
        self.app.addButton("Трапеция", lambda: self.calculate("Трапеция"))
        self.app.addLabel("result", "", colspan=2)
        self.app.addButton("Сохранить", self.save)
        self.app.go()

    def calculate(self, shape_type):
        try:
            if shape_type == "Прямоугольник":
                a = float(self.app.stringBox("Ввод", "Сторона a:"))
                b = float(self.app.stringBox("Ввод", "Сторона b:"))
                shape = Rectangle(a, b)
            elif shape_type == "Треугольник":
                a = float(self.app.stringBox("Ввод", "Сторона a:"))
                b = float(self.app.stringBox("Ввод", "Сторона b:"))
                c = float(self.app.stringBox("Ввод", "Сторона c:"))
                shape = Triangle(a, b, c)
            elif shape_type == "Трапеция":
                a = float(self.app.stringBox("Ввод", "Основание a:"))
                b = float(self.app.stringBox("Ввод", "Основание b:"))
                c = float(self.app.stringBox("Ввод", "Боковая сторона c:"))
                d = float(self.app.stringBox("Ввод", "Боковая сторона d:"))
                shape = Trapezoid(a, b, c, d)

            result_text = f"{shape}\nПлощадь: {shape.area():.2f}\n"
            result_text += f"Радиус описанной окружности: {shape.circumradius():.2f}\n"
            result_text += f"Радиус вписанной окружности: {shape.inradius():.2f}"

            self.app.setLabel("result", result_text)
            self.last_result = result_text

        except ValueError:
            self.app.errorBox("Ошибка", "Некорректный ввод!")
        except Exception as e:
            self.app.errorBox("Ошибка", f"Ошибка: {str(e)}")

    def save(self):
        if not hasattr(self, 'last_result'):
            self.app.errorBox("Ошибка", "Нет данных для сохранения")
            return

        doc = Document()
        doc.add_paragraph(self.last_result)
        doc.save("result.docx")
        self.app.infoBox("Успех", "Сохранено в result.docx")

if __name__ == "__main__":
    SimpleGeometryCalculator()
